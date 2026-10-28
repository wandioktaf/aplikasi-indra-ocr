import streamlit as strl
import indra_ocr
import os
import re
from docx import Document
from datetime import datetime

# 1. Konfigurasi Halaman Web
strl.set_page_config(page_title="Indra OCR Super App", page_icon="🚀", layout="wide")

# Inisialisasi Riwayat Proses (Session State) jika belum ada
if "riwayat_proses" not in strl.session_state:
    strl.session_state.riwayat_proses = []

# 2. Pembuatan Sidebar untuk Fitur Kontrol (Pengaturan CPU & Riwayat)
strl.sidebar.title("⚙️ Pengaturan & Kontrol")

# Fitur Pengaturan Core CPU (Threads)
slider_threads = strl.sidebar.slider(
    "Jumlah Core CPU (Threads):", 
    min_value=1, 
    max_value=8, 
    value=4,
    help="Semakin tinggi angkanya, proses OCR dokumen besar akan semakin cepat."
)

# Fungsi untuk memuat model AI (di-cache berdasarkan jumlah threads yang dipilih)
@strl.cache_resource
def muat_model_ocr(threads):
    return indra_ocr.load(threads=threads)

try:
    pipeline = muat_model_ocr(slider_threads)
except Exception as e:
    strl.sidebar.error(f"Gagal memuat model AI: {e}")
    strl.stop()

# Menampilkan Riwayat Proses di Sidebar
strl.sidebar.markdown("---")
strl.sidebar.subheader("⏳ Riwayat Proses Sesi Ini")
if not strl.session_state.riwayat_proses:
    strl.sidebar.info("Belum ada file yang diproses.")
else:
    for item in strl.session_state.riwayat_proses:
        strl.sidebar.write(f"⏱️ {item['waktu']} - **{item['nama']}** ({item['status']})")

# 3. Konten Utama Aplikasi Web
strl.title("🚀 Indra OCR Super App (GUI Version)")
strl.write("Aplikasi ekstraksi dokumen canggih dengan fitur pembersih teks, pencarian kata, dan ekspor Word.")

# Komponen Drag-and-Drop File Uploader
file_diunggah = strl.file_uploader("Pilih atau seret dokumen Anda ke sini (Mendukung file besar)", type=["pdf", "png", "jpg", "jpeg"])

# Fungsi untuk membersihkan format Markdown menjadi teks polos biasa (Plain Text)
def bersihkan_ke_teks_polos(teks_md):
    teks = re.sub(r'#+\s+', '', teks_md) # Hapus header markdown (#)
    teks = re.sub(r'[*_~`]', '', teks)   # Hapus format tebal/miring (*, _, ~, `)
    teks = re.sub(r'\[(.*?)\]\(.*?\)', r'\1', teks) # Hapus link, sisakan teksnya
    teks = re.sub(r'^-\s+', '', teks, flags=re.MULTILINE) # Hapus tanda bullet list (-)
    teks = re.sub(r'\|\s*', ' ', teks)    # Hapus garis pembatas tabel (|)
    return teks

# Fungsi untuk membuat file Microsoft Word (.docx) secara langsung
def buat_file_word(teks_konten):
    doc = Document()
    doc.add_heading('Hasil Ekstraksi Teks - Indra OCR', level=1)
    doc.add_paragraph(f"Diproses pada: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
    doc.add_paragraph(teks_konten)
    
    path_word = "hasil_ocr_sementara.docx"
    doc.save(path_word)
    with open(path_word, "rb") as f:
        byte_data = f.read()
    os.remove(path_word) # Hapus file lokal setelah dibaca ke memori
    return byte_data

if file_diunggah is not None:
    strl.success(f"File siap diproses: **{file_diunggah.name}**")
    
    # Tombol Eksekusi Utama
    if strl.button("Mulai Ekstrak Teks ✨", type="primary"):
        path_sementara = os.path.join(".", file_diunggah.name)
        with open(path_sementara, "wb") as f:
            f.write(file_diunggah.getbuffer())
            
        teks_mentah_ocr = ""
        waktu_sekarang = datetime.now().strftime("%H:%M:%S")
        
        with strl.spinner("Model AI sedang memindai dan membaca dokumen Anda... Mohon tunggu..."):
            try:
                for index_halaman, halaman in pipeline.convert(path_sementara):
                    teks_mentah_ocr += f"=== HALAMAN {index_halaman + 1} ===\n{halaman.markdown}\n\n"
                
                strl.balloons()
                
                # Simpan data ke session state agar data teks tidak hilang saat tombol UI lain diklik
                strl.session_state.teks_hasil = teks_mentah_ocr
                strl.session_state.nama_file_aktif = file_diunggah.name
                
                # Tambahkan ke riwayat sukses
                strl.session_state.riwayat_proses.append({"waktu": waktu_sekarang, "nama": file_diunggah.name, "status": "Sukses"})
                
            except Exception as e:
                strl.error(f"Terjadi kesalahan: {e}")
                strl.session_state.riwayat_proses.append({"waktu": waktu_sekarang, "nama": file_diunggah.name, "status": "Gagal"})
            finally:
                if os.path.exists(path_sementara):
                    os.remove(path_sementara)

    # Menampilkan panel hasil jika data teks sudah ada di memori aplikasi
    if "teks_hasil" in strl.session_state and strl.session_state.nama_file_aktif == file_diunggah.name:
        strl.markdown("---")
        strl.header("📝 Panel Hasil Ekstraksi & Fitur Tambahan")
        
        # 1. Fitur Pembersih Teks Otomatis (Toggle Pilihan Tampilan)
        opsi_tampilan = strl.radio(
            "Pilih Format Tampilan Teks:",
            ("Format Asli AI (Markdown / Tabel)", "Format Bersih (Teks Polos / Plain Text)"),
            horizontal=True
        )
        
        if opsi_tampilan == "Format Bersih (Teks Polos / Plain Text)":
            teks_final = bersihkan_ke_teks_polos(strl.session_state.teks_hasil)
        else:
            teks_final = strl.session_state.teks_hasil

        # 2. Fitur Kotak Pencarian Kata (Search Text)
        kata_kunci = strl.text_input("🔍 Cari kata kunci tertentu di dalam teks hasil:")
        if kata_kunci:
            # Cari berapa kali kata kunci muncul (Case-Insensitive)
            kemunculan = len(re.findall(re.escape(kata_kunci), teks_final, re.IGNORECASE))
            if kemunculan > 0:
                strl.success(f"Kata '{kata_kunci}' ditemukan sebanyak **{kemunculan} kali** di dalam dokumen!")
            else:
                strl.warning(f"Kata '{kata_kunci}' tidak ditemukan di dalam dokumen.")

        # Menampilkan Teks Hasil Pemuatan
        strl.code(teks_final, language="markdown", line_numbers=True)
        
        # 3. Fitur Multi-Format Download (Ekspor ke TXT & Word .DOCX)
        kolom1, kolom2 = strl.columns(2)
        
        with kolom1:
            strl.download_button(
                label="📥 Unduh sebagai File Teks (.txt)",
                data=teks_final,
                file_name=f"hasil_{os.path.splitext(file_diunggah.name)[0]}.txt",
                mime="text/plain",
                use_container_width=True
            )
            
        with kolom2:
            data_word = buat_file_word(teks_final)
            strl.download_button(
                label="📥 Unduh sebagai Dokumen Word (.docx)",
                data=data_word,
                file_name=f"hasil_{os.path.splitext(file_diunggah.name)[0]}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
